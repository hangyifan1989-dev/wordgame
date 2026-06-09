import openpyxl
import json
from pathlib import Path

BOOK1 = Path(__file__).parent / "新1单词表.xlsx"
BOOK2 = Path(__file__).parent / "新2单词表.xlsx"

class Word:
    __slots__ = ('lesson', 'seq', 'word', 'phonetic', 'pos', 'chinese')
    def __init__(self, lesson, seq, word, phonetic, pos, chinese):
        self.lesson = lesson
        self.seq = seq
        self.word = word
        self.phonetic = phonetic
        self.pos = pos
        self.chinese = chinese

    def to_dict(self):
        return {
            'lesson': self.lesson,
            'seq': self.seq,
            'word': self.word,
            'phonetic': self.phonetic,
            'pos': self.pos,
            'chinese': self.chinese,
        }

def _load_book1():
    wb = openpyxl.load_workbook(BOOK1)
    ws = wb.active
    words = {}
    current_lesson = 0
    for row in range(3, ws.max_row + 1):
        c2 = ws.cell(row, 2).value
        c3 = ws.cell(row, 3).value
        if c2 is not None and str(c2).startswith('Lesson'):
            current_lesson = int(str(c2).replace('Lesson', ''))
            if c3 is None:
                continue
            seq = int(c3)
            word = str(ws.cell(row, 5).value or '').strip()
            phonetic = str(ws.cell(row, 6).value or '').strip()
            pos = str(ws.cell(row, 7).value or '').strip()
            chinese = str(ws.cell(row, 8).value or '').strip()
        elif c2 is None and c3 is not None and current_lesson > 0:
            seq = int(c3)
            word = str(ws.cell(row, 5).value or '').strip()
            phonetic = str(ws.cell(row, 6).value or '').strip()
            pos = str(ws.cell(row, 7).value or '').strip()
            chinese = str(ws.cell(row, 8).value or '').strip()
        else:
            continue
        if word and word != 'None':
            key = (current_lesson, seq)
            words[key] = Word(current_lesson, seq, word, phonetic, pos, chinese)
    wb.close()
    return words

def _load_book2():
    wb = openpyxl.load_workbook(BOOK2)
    ws = wb.active
    words = {}
    current_lesson = 0
    for row in range(3, ws.max_row + 1):
        c2 = ws.cell(row, 2).value
        c3 = ws.cell(row, 3).value
        if c2 is not None and str(c2).startswith('Lesson'):
            current_lesson = int(str(c2).replace('Lesson', ''))
            if c3 is None:
                continue
            seq = int(c3)
            word = str(ws.cell(row, 4).value or '').strip()
            phonetic = str(ws.cell(row, 5).value or '').strip()
            pos = str(ws.cell(row, 6).value or '').strip()
            chinese = str(ws.cell(row, 7).value or '').strip()
        elif c2 is None and c3 is not None and current_lesson > 0:
            seq = int(c3)
            word = str(ws.cell(row, 4).value or '').strip()
            phonetic = str(ws.cell(row, 5).value or '').strip()
            pos = str(ws.cell(row, 6).value or '').strip()
            chinese = str(ws.cell(row, 7).value or '').strip()
        else:
            continue
        if word and word != 'None':
            key = (current_lesson, seq)
            words[key] = Word(current_lesson, seq, word, phonetic, pos, chinese)
    wb.close()
    return words

_cache = {}

def load_book(book_id):
    if book_id in _cache:
        return _cache[book_id]
    if book_id == 1:
        raw = _load_book1()
    else:
        raw = _load_book2()
    lessons = {}
    for (lesson_num, seq), w in raw.items():
        lessons.setdefault(lesson_num, []).append(w)
    for lnum in sorted(lessons.keys()):
        lessons[lnum].sort(key=lambda w: w.seq)
    _cache[book_id] = lessons
    return lessons

def get_lesson_words(book_id, lesson_num):
    lessons = load_book(book_id)
    return lessons.get(lesson_num, [])

def get_total_lessons(book_id):
    lessons = load_book(book_id)
    return len(lessons)

def get_all_lesson_numbers(book_id):
    lessons = load_book(book_id)
    return sorted(lessons.keys())

def get_score_per_word(book_id, lesson_num):
    if book_id == 1:
        return 1
    lesson_idx = lesson_num
    tier = (lesson_idx - 1) // 10
    return 2 + tier
